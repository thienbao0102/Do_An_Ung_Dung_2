from __future__ import annotations

import io
import math
import re
import textwrap
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


SOURCE = Path(r"F:\Downloads\BM-29-NCKHSV-bao-cao-tong-ket-de-tai.docx")
WORK = Path(r"F:\Github\Do_An_Ung_Dung_2\.codex-work")
ASSETS = WORK / "report_assets"
OUTPUT = Path(r"F:\Github\Do_An_Ung_Dung_2\BM-29-ScoreSense-bao-cao-tong-ket-da-bo-sung.docx")
BULLET_NUM_IDS = {}

ASSETS.mkdir(parents=True, exist_ok=True)


def font(path: str, size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts") / path,
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def draw_wrapped_center(draw: ImageDraw.ImageDraw, box, text, fnt, fill, spacing=8):
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 36
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = (current + " " + word).strip()
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    heights = [draw.textbbox((0, 0), line, font=fnt)[3] for line in lines]
    total = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=fnt)
        width = bbox[2] - bbox[0]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=fnt, fill=fill)
        y += height + spacing


def arrow(draw, start, end, color=(57, 90, 140), width=7):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 20
    for offset in (2.55, -2.55):
        point = (
            end[0] + length * math.cos(angle + offset),
            end[1] + length * math.sin(angle + offset),
        )
        draw.line([end, point], fill=color, width=width)


def rounded_box(draw, box, title, subtitle, fill, outline=(42, 73, 115)):
    draw.rounded_rectangle(box, radius=24, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = box
    draw_wrapped_center(draw, (x1, y1 + 18, x2, y1 + 78), title, font("arialbd.ttf", 27, True), (20, 48, 80))
    draw_wrapped_center(draw, (x1, y1 + 80, x2, y2 - 10), subtitle, font("arial.ttf", 22), (35, 55, 75), spacing=5)


def make_architecture_diagrams():
    canvas = Image.new("RGB", (1800, 950), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((70, 35), "KIẾN TRÚC PHẦN MỀM SCORESENSE", font=font("arialbd.ttf", 38, True), fill=(22, 55, 94))
    boxes = [
        (70, 180, 400, 650),
        (520, 180, 850, 650),
        (970, 180, 1300, 650),
        (1420, 180, 1750, 650),
    ]
    rounded_box(draw, boxes[0], "LỚP GIAO DIỆN", "Flutter Web\nBiểu mẫu nhiều bước\nNhập tệp CSV\nHiển thị Pass/Fail và bảng kết quả", (225, 240, 255))
    rounded_box(draw, boxes[1], "LỚP GIAO TIẾP", "HTTP/JSON\nDart package http\nAPI_BASE_URL cấu hình khi build\nTimeout 120 giây", (231, 247, 239))
    rounded_box(draw, boxes[2], "LỚP DỊCH VỤ", "Flask + Flask-CORS\n/health, /predict, /totalmodel\nPandas tiền xử lý\nGiới hạn request 10 MB", (255, 244, 220))
    rounded_box(draw, boxes[3], "LỚP MÔ HÌNH", "DecisionTreeRegressor (.pkl)\nDanh sách 41 cột (.json)\nBa phiên bản đóng gói\nNạp và cache khi suy luận", (244, 230, 250))
    for a, b in zip(boxes, boxes[1:]):
        arrow(draw, (a[2] + 15, 415), (b[0] - 15, 415))
    draw.rounded_rectangle((160, 760, 1640, 885), radius=22, fill=(245, 247, 250), outline=(90, 110, 135), width=3)
    draw_wrapped_center(draw, (180, 770, 1620, 875), "Triển khai: Flutter Web trên Firebase Hosting; Flask API chạy bằng Gunicorn trên Render. Hệ thống hiện không có cơ sở dữ liệu, xác thực người dùng hoặc lưu lịch sử dự đoán.", font("arial.ttf", 25), (35, 50, 70), spacing=6)
    software = ASSETS / "kien_truc_phan_mem.png"
    canvas.save(software)

    canvas = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((70, 35), "KIẾN TRÚC AI VÀ LUỒNG SUY LUẬN", font=font("arialbd.ttf", 38, True), fill=(22, 55, 94))
    train_box = (90, 160, 800, 420)
    rounded_box(draw, train_box, "NHÁNH HUẤN LUYỆN NGOẠI TUYẾN", "student_data.csv → get_dummies(drop_first=True) → tách G3 → train/test 80/20 → DecisionTreeRegressor(random_state=42) → lưu model .pkl và train_columns .json", (231, 247, 239))
    online_boxes = [
        (90, 600, 350, 900),
        (440, 600, 700, 900),
        (790, 600, 1050, 900),
        (1140, 600, 1400, 900),
        (1490, 600, 1750, 900),
    ]
    labels = [
        ("DỮ LIỆU MỚI", "Biểu mẫu một SV hoặc ma trận từ CSV"),
        ("TIỀN XỬ LÝ", "Bỏ cột name; one-hot encoding"),
        ("ĐỒNG BỘ", "Reindex theo 41 cột; cột thiếu điền 0"),
        ("SUY LUẬN", "Nạp model theo version; predict điểm G3"),
        ("PHẢN HỒI", "JSON: name, prediction, originalIndex; UI dùng >10 để gán Pass"),
    ]
    colors = [(225, 240, 255), (231, 247, 239), (255, 244, 220), (244, 230, 250), (255, 233, 233)]
    for box, (title, subtitle), color in zip(online_boxes, labels, colors):
        rounded_box(draw, box, title, subtitle, color)
    for a, b in zip(online_boxes, online_boxes[1:]):
        arrow(draw, (a[2] + 12, 750), (b[0] - 12, 750))
    arrow(draw, (650, train_box[3] + 10), (1270, online_boxes[3][1] - 15), color=(80, 105, 130), width=6)
    draw.text((720, 475), "Cung cấp artifact mô hình và schema đặc trưng", font=font("arial.ttf", 23), fill=(50, 70, 90))
    ai = ASSETS / "kien_truc_ai.png"
    canvas.save(ai)
    return software, ai


def extract_original_images():
    files = []
    with zipfile.ZipFile(SOURCE) as archive:
        for index in range(1, 5):
            name = f"word/media/image{index}.png"
            target = ASSETS / f"giao_dien_{index}.png"
            target.write_bytes(archive.read(name))
            files.append(target)
    return files


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is not None:
        return
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def finalize_accessibility(doc):
    picture_descriptions = iter([
        "Sơ đồ kiến trúc phần mềm ScoreSense",
        "Sơ đồ kiến trúc AI và luồng suy luận ScoreSense",
        "Giao diện lựa chọn phương thức dự đoán",
        "Giao diện biểu mẫu nhập dữ liệu cá nhân",
        "Giao diện kết quả dự đoán cá nhân",
        "Giao diện kết quả dự đoán hàng loạt",
    ])
    for doc_pr in doc._element.iter(qn("wp:docPr")):
        if doc_pr.get("descr") or doc_pr.get("title"):
            continue
        name = doc_pr.get("name", "")
        if name.startswith("Rectangle"):
            description = "Khung viền trang bìa báo cáo"
        else:
            description = next(picture_descriptions, "Hình minh họa trong báo cáo ScoreSense")
        doc_pr.set("descr", description)
        doc_pr.set("title", description)

    # The source template contains two layout tables on the cover pages; marking
    # their first rows also lets assistive technology traverse every table cleanly.
    for table_element in doc._element.iter(qn("w:tbl")):
        first_row = table_element.find(qn("w:tr"))
        if first_row is None:
            continue
        row_properties = first_row.find(qn("w:trPr"))
        if row_properties is None:
            row_properties = OxmlElement("w:trPr")
            first_row.insert(0, row_properties)
        if row_properties.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row_properties.append(header)


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(Cm(width).twips)))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(int(Cm(width).twips) for width in widths_cm)))
    tbl_w.set(qn("w:type"), "dxa")


def format_run(run, size=13, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def style_body_paragraph(paragraph, first_line=True):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(1.0) if first_line else None
    for run in paragraph.runs:
        format_run(run)


def add_body(doc, text, first_line=True, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        format_run(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        format_run(r)
    else:
        r = p.add_run(text)
        format_run(r)
    style_body_paragraph(p, first_line=first_line)
    return p


def add_heading(doc, text, level, page_break=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level > 1 else 12)
    p.paragraph_format.space_after = Pt(5)
    if page_break:
        p.paragraph_format.page_break_before = True
    for run in p.runs:
        format_run(run, size=13 if level > 1 else 14, bold=True)
    return p


def add_opening_heading(doc, text):
    p = doc.add_paragraph(text, style="Opening Heading")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        format_run(run, size=13, bold=True)
    return p


def add_title(doc, text):
    p = doc.add_paragraph(text, style="Unnumbered Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(12)
    for run in p.runs:
        format_run(run, size=14, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Paragraph")
    p.add_run(text)
    numbering = doc.part.numbering_part.element
    key = id(doc)
    if key not in BULLET_NUM_IDS:
        abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
        num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
        abstract_id = max(abstract_ids, default=0) + 1
        num_id = max(num_ids, default=0) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        level.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
        abstract.append(level)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        BULLET_NUM_IDS[key] = num_id
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(BULLET_NUM_IDS[key]))
    num_pr.extend([ilvl, num_id_node])
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        format_run(run)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    match = re.match(r"^(Bảng|Hình)\s+(\d+)\s*:\s*(.+)$", text)
    if match:
        label, display_number, title = match.groups()
        lead = p.add_run(f"{label} ")
        format_run(lead, size=11, italic=True)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" SEQ {label} \\* ARABIC "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        result = OxmlElement("w:t")
        result.text = display_number
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        field_run = p.add_run()
        field_run._r.extend([begin, instr, separate, result, end])
        format_run(field_run, size=11, italic=True)
        tail = p.add_run(f": {title}")
        format_run(tail, size=11, italic=True)
    else:
        run = p.add_run(text)
        format_run(run, size=11, italic=True)
    return p


def add_picture(doc, path, width_cm=15.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def add_table(doc, headers, rows, widths_cm, font_size=10.5, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        format_run(r, size=font_size, bold=True)
        set_cell_margins(cell)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = aligns[index] if aligns else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            format_run(r, size=font_size)
            set_cell_margins(cell)
    set_table_widths(table, widths_cm)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def remove_old_body(doc):
    starts = [p for p in doc.paragraphs if p.text.strip() == "MỞ ĐẦU" and p.style.name == "Title"]
    if not starts:
        raise RuntimeError("Không tìm thấy điểm bắt đầu phần nội dung cũ.")
    start = starts[-1]._p
    body = doc.element.body
    removing = False
    for child in list(body):
        if child is start:
            removing = True
        if removing and child.tag != qn("w:sectPr"):
            body.remove(child)


def set_document_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(13)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.bold = True
    if "Opening Heading" not in [style.name for style in doc.styles]:
        opening = doc.styles.add_style("Opening Heading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        opening = doc.styles["Opening Heading"]
    opening.base_style = doc.styles["Normal"]
    opening.font.name = "Times New Roman"
    opening._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    opening._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    opening.font.size = Pt(13)
    opening.font.bold = True
    p_pr = opening._element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), "1")
    if "Unnumbered Heading 1" not in [style.name for style in doc.styles]:
        unnumbered = doc.styles.add_style("Unnumbered Heading 1", WD_STYLE_TYPE.PARAGRAPH)
    else:
        unnumbered = doc.styles["Unnumbered Heading 1"]
    unnumbered.base_style = doc.styles["Normal"]
    unnumbered.font.name = "Times New Roman"
    unnumbered._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    unnumbered._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    unnumbered.font.size = Pt(14)
    unnumbered.font.bold = True
    unnumbered_p_pr = unnumbered._element.get_or_add_pPr()
    unnumbered_outline = unnumbered_p_pr.find(qn("w:outlineLvl"))
    if unnumbered_outline is None:
        unnumbered_outline = OxmlElement("w:outlineLvl")
        unnumbered_p_pr.append(unnumbered_outline)
    unnumbered_outline.set(qn("w:val"), "0")


def build_report():
    software_diagram, ai_diagram = make_architecture_diagrams()
    screenshots = extract_original_images()
    doc = Document(SOURCE)
    remove_old_body(doc)
    set_document_styles(doc)

    add_title(doc, "MỞ ĐẦU")
    add_opening_heading(doc, "Tính cấp thiết của đề tài")
    add_body(doc, "Các chỉ báo học vụ như điểm quá trình, điểm giữa kỳ và số buổi vắng học chứa thông tin quan trọng về nguy cơ suy giảm kết quả. Khi được khai thác bằng học máy, các chỉ báo này có thể hỗ trợ sinh viên và người làm công tác cố vấn nhận biết sớm xu hướng học tập. Giá trị của hệ thống dự báo không nằm ở việc thay thế đánh giá chuyên môn, mà ở khả năng cung cấp thêm một tín hiệu định lượng, kịp thời và có thể kiểm chứng.")
    add_body(doc, "ScoreSense được xây dựng như một nguyên mẫu phần mềm hoàn chỉnh gồm giao diện Flutter Web, dịch vụ Flask và mô hình Decision Tree Regressor. Việc kết hợp thành phần AI với một luồng sử dụng thực tế giúp đề tài vượt khỏi phạm vi thử nghiệm thuật toán đơn lẻ. Tuy nhiên, để báo cáo phản ánh trung thực sản phẩm, phạm vi dữ liệu, kiến trúc triển khai, giới hạn của mô hình và trạng thái kiểm thử phải được mô tả đúng theo mã nguồn hiện có.")
    add_opening_heading(doc, "Mục tiêu nghiên cứu")
    add_body(doc, "Mục tiêu chung của đề tài là xây dựng và đánh giá một nguyên mẫu hệ thống dự đoán điểm cuối kỳ G3, cho phép nhập dữ liệu của một người học hoặc tải tệp CSV để dự đoán hàng loạt, qua đó kiểm chứng khả năng tích hợp mô hình cây quyết định vào ứng dụng web.")
    for item in [
        "Tổ chức và tiền xử lý dữ liệu theo đúng lược đồ 32 biến đầu vào của bộ Student Performance.",
        "Huấn luyện Decision Tree Regressor, lưu mô hình cùng danh sách cột sau mã hóa và cung cấp suy luận theo phiên bản.",
        "Xây dựng giao diện Flutter Web và API Flask hỗ trợ dự đoán cá nhân, dự đoán hàng loạt và hiển thị kết quả Pass/Fail.",
        "So sánh mô hình hiện tại với một số thuật toán hồi quy trên cùng cách chia dữ liệu để làm rõ đánh đổi giữa sai số, độ phức tạp và khả năng diễn giải.",
        "Kiểm thử các hợp đồng API, kiểm thử widget hiện có và phân tích tĩnh mã nguồn để xác định mức độ sẵn sàng của sản phẩm.",
    ]:
        add_bullet(doc, item)
    add_opening_heading(doc, "Đối tượng và phạm vi nghiên cứu")
    add_body(doc, "Đối tượng kỹ thuật của nghiên cứu là quy trình dự đoán G3 từ dữ liệu nhân khẩu học, gia đình, hành vi học tập và hai điểm giai đoạn G1, G2. Dữ liệu trong kho mã nguồn gồm 395 bản ghi, 33 cột và không có giá trị thiếu. Đây là phần Toán học của bộ Student Performance do Cortez và Silva công bố, thu thập tại hai trường trung học ở Bồ Đào Nha; không phải dữ liệu khảo sát sinh viên Trường Đại học Văn Lang [1], [2].")
    add_body(doc, "Vì vậy, kết quả trong báo cáo được giới hạn ở mức kiểm chứng kỹ thuật và nguyên mẫu ứng dụng. Mô hình chưa đủ cơ sở để suy rộng trực tiếp cho sinh viên Việt Nam hoặc sử dụng như công cụ ra quyết định học vụ. Trước khi triển khai tại Văn Lang, hệ thống cần được huấn luyện và đánh giá lại trên dữ liệu nội bộ có đồng thuận, cơ chế bảo vệ dữ liệu và quy trình giám sát phù hợp.")
    add_opening_heading(doc, "Cấu trúc báo cáo")
    add_body(doc, "Báo cáo gồm năm chương: Chương 1 xác lập bài toán và phạm vi; Chương 2 trình bày dữ liệu, phương pháp và so sánh mô hình; Chương 3 mô tả kiến trúc phần mềm và kiến trúc AI; Chương 4 báo cáo kết quả thực nghiệm, tích hợp và kiểm thử; Chương 5 tổng kết các đóng góp, giới hạn và hướng phát triển.")

    add_heading(doc, "CHƯƠNG 1. TỔNG QUAN VÀ XÁC LẬP BÀI TOÁN", 1, page_break=True)
    add_heading(doc, "Bài toán nghiệp vụ", 2)
    add_body(doc, "ScoreSense tiếp nhận các thuộc tính của người học và dự đoán điểm G3 trên thang 0-20. Ở lớp giao diện, điểm dự đoán được ánh xạ thành hai trạng thái: Pass khi giá trị lớn hơn 10 và Fail khi giá trị nhỏ hơn hoặc bằng 10. Hệ thống phục vụ hai tình huống: một người dùng nhập biểu mẫu nhiều bước; hoặc người dùng tải tệp CSV để nhận kết quả cho nhiều bản ghi kèm thứ tự dòng ban đầu.")
    add_body(doc, "Bài toán cốt lõi là hồi quy, không phải phân loại. Decision Tree Regressor phù hợp với kiểu dữ liệu mục tiêu là điểm số; nhãn Pass/Fail chỉ là quy tắc trình bày ở giao diện sau suy luận. Việc phân biệt hai tầng này giúp lựa chọn đúng chỉ số đánh giá: MSE, RMSE, MAE và R² cho mô hình hồi quy; tỷ lệ Pass/Fail chỉ dùng cho trực quan hóa kết quả.")
    add_heading(doc, "Nguồn dữ liệu và mức độ phù hợp", 2)
    add_body(doc, "Bộ dữ liệu gốc có 395 bản ghi của môn Toán, 32 biến dự báo và biến đích G3. Nhóm biến gồm thông tin cá nhân (school, sex, age, address), bối cảnh gia đình (famsize, Pstatus, Medu, Fedu, Mjob, Fjob), điều kiện và hành vi học tập (traveltime, studytime, failures, hỗ trợ học tập, hoạt động, Internet, absences) cùng hai điểm G1 và G2 [1], [2].")
    add_body(doc, "Ưu điểm của bộ dữ liệu là lược đồ rõ ràng, không có giá trị thiếu và đủ nhỏ để kiểm chứng toàn bộ chuỗi xử lý. Hạn chế quan trọng là khác biệt về cấp học, quốc gia và bối cảnh quản trị. Do đó, dữ liệu phù hợp để chứng minh kiến trúc và quy trình kỹ thuật, nhưng không phù hợp để khẳng định hiệu lực dự báo cho sinh viên Văn Lang.")
    add_heading(doc, "Tiêu chí lựa chọn mô hình", 2)
    add_body(doc, "Mô hình của đề tài được đánh giá theo bốn tiêu chí: sai số trên dữ liệu chưa thấy; khả năng giải thích cho người không chuyên; chi phí huấn luyện và suy luận; mức độ thuận tiện khi tích hợp vào backend. Cây hồi quy tạo các phân vùng theo điều kiện thuộc tính và dự đoán giá trị trung bình tại nút lá, phù hợp với bài toán G3 và có cấu trúc trực quan hơn các mô hình tổ hợp [3].")
    add_caption(doc, "Bảng 1: So sánh định tính các nhóm mô hình được xem xét")
    add_table(doc,
        ["Mô hình", "Ưu điểm", "Hạn chế", "Mức phù hợp với ScoreSense"],
        [
            ["Hồi quy tuyến tính", "Nhanh, hệ số dễ kiểm tra", "Khó mô tả quan hệ phi tuyến và tương tác", "Dùng làm đường cơ sở"],
            ["KNN Regressor", "Đơn giản, không giả định dạng hàm", "Phụ thuộc khoảng cách; suy luận tăng theo dữ liệu", "Phù hợp đối chứng"],
            ["Decision Tree Regressor", "Phi tuyến, suy luận nhanh, có thể truy vết nhánh", "Nhạy dữ liệu; dễ quá khớp khi cây sâu", "Mô hình triển khai hiện tại"],
            ["Random Forest Regressor", "Ổn định và thường giảm sai số", "Khó giải thích từng dự đoán; mô hình lớn hơn", "Ứng viên nâng cấp khi ưu tiên hiệu năng [4]"],
            ["Gradient Boosting Regressor", "Hiệu năng tốt trên dữ liệu bảng", "Tối ưu và diễn giải phức tạp hơn", "Ứng viên thử nghiệm bổ sung"],
        ],
        [3.3, 4.2, 4.2, 4.3], font_size=9.5)
    add_heading(doc, "Phạm vi đóng góp của đề tài", 2)
    add_body(doc, "Đóng góp chính của đề tài là tích hợp được một quy trình dự đoán đầu-cuối: chuẩn bị đặc trưng, lưu artifact mô hình, cung cấp API và xây dựng giao diện sử dụng. Đề tài chưa tạo ra bộ dữ liệu sinh viên mới, chưa triển khai cơ chế cố vấn tự động và chưa chứng minh tác động giáo dục. Cách xác định phạm vi này giúp tách rõ kết quả đã thực hiện khỏi các định hướng phát triển.")

    add_heading(doc, "CHƯƠNG 2. PHƯƠNG PHÁP VÀ THIẾT KẾ THỰC NGHIỆM", 1, page_break=True)
    add_heading(doc, "Quy trình nghiên cứu bám theo mã nguồn", 2)
    add_body(doc, "Quy trình được tái lập từ các tệp student_data.csv, decisiontree.py, ba tệp mô hình finalized_model_v*.pkl, ba danh sách train_columns*.json và lớp xử lý handelApi.py. Các bước gồm: đọc dữ liệu; one-hot encoding với drop_first=True; tách G3; chia tập 80/20 với random_state=42; huấn luyện Decision Tree Regressor; lưu mô hình và danh sách cột; nạp artifact trong API để suy luận.")
    add_heading(doc, "Tiền xử lý dữ liệu", 2)
    add_body(doc, "Các biến phân loại được chuyển sang biến giả bằng pandas.get_dummies. Sau mã hóa, 32 biến đầu vào tạo thành 41 đặc trưng. Khi suy luận, dữ liệu mới được mã hóa độc lập rồi reindex theo danh sách 41 cột đã lưu; cột không xuất hiện được điền 0. Cột name, nếu có, được tách trước suy luận và ghép lại vào kết quả. Cách làm này bảo đảm số lượng và thứ tự đặc trưng tương thích với mô hình.")
    add_caption(doc, "Bảng 2: Thông số dữ liệu và thực nghiệm hiện tại")
    add_table(doc,
        ["Thuộc tính", "Giá trị xác minh từ dự án"],
        [
            ["Số bản ghi / số cột", "395 / 33"],
            ["Biến đầu vào / biến mục tiêu", "32 biến / G3 (0-20)"],
            ["Số đặc trưng sau mã hóa", "41"],
            ["Giá trị thiếu", "0"],
            ["Chia dữ liệu", "80% huấn luyện (316), 20% kiểm thử (79)"],
            ["Khả năng tái lập", "random_state = 42"],
        ], [6.1, 9.9], font_size=10.5)
    add_heading(doc, "Cấu hình mô hình đang triển khai", 2)
    add_body(doc, "Ba artifact mô hình trong thư mục backend_scoresense/models đều là DecisionTreeRegressor. Cấu hình chính gồm criterion='squared_error', splitter='best', random_state=42, không giới hạn max_depth và min_samples_leaf=1. Cây hiện có độ sâu 11 và 127 nút lá. Cấu hình này giữ nguyên hành vi của dự án, nhưng cho thấy mô hình chưa được cắt tỉa; vì vậy khả năng giải thích chỉ nên được xem là tương đối và cần truy vết theo đường đi cụ thể thay vì trình bày toàn bộ cây.")
    add_heading(doc, "Thiết kế so sánh mô hình", 2)
    add_body(doc, "Để trả lời yêu cầu lựa chọn mô hình, năm thuật toán được huấn luyện lại trên cùng dữ liệu đã mã hóa, cùng phép chia 80/20 và cùng random_state khi thuật toán hỗ trợ. So sánh dùng MAE, RMSE và R²; giá trị thấp hơn là tốt hơn đối với MAE/RMSE, còn R² cao hơn là tốt hơn. Đây là đối chứng trên một lần chia dữ liệu, không thay thế đánh giá chéo nhiều lần.")
    add_caption(doc, "Bảng 3: Kết quả so sánh mô hình trên cùng tập kiểm thử 79 bản ghi")
    add_table(doc,
        ["Mô hình", "MAE", "RMSE", "R²", "Nhận định"],
        [
            ["Linear Regression", "1,6467", "2,3784", "0,7241", "Đường cơ sở, dễ diễn giải hệ số"],
            ["KNN (k=5)", "1,4304", "2,2243", "0,7587", "Sai số giảm nhưng suy luận phụ thuộc dữ liệu"],
            ["Decision Tree", "1,1392", "2,0500", "0,7950", "MAE thấp nhất trong so sánh; mô hình hiện tại"],
            ["Random Forest (100 cây)", "1,1646", "1,9488", "0,8148", "RMSE/R² tốt nhất; khó giải thích hơn"],
            ["Gradient Boosting", "1,1594", "2,0047", "0,8040", "Hiệu năng tốt, quy trình phức tạp hơn"],
        ], [3.8, 2.4, 2.4, 2.4, 5.0], font_size=9.2,
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_heading(doc, "Lý do sử dụng Decision Tree Regressor", 2)
    for item in [
        "Đúng loại bài toán: biến G3 là điểm số, vì vậy hồi quy phù hợp hơn mô hình phân loại.",
        "Hiệu năng đủ cạnh tranh: Decision Tree đạt MAE 1,1392, thấp nhất trong nhóm đối chứng; chênh lệch R² với Random Forest chỉ khoảng 0,0198 trên phép chia hiện tại.",
        "Khả năng truy vết: một dự đoán có thể được giải thích bằng chuỗi điều kiện từ nút gốc đến nút lá; mô hình tổ hợp cần thêm kỹ thuật giải thích hậu nghiệm.",
        "Chi phí triển khai thấp: artifact khoảng 20 KB, nạp nhanh và suy luận trực tiếp trong Flask mà không cần dịch vụ AI riêng.",
        "Phù hợp quy mô dữ liệu: 395 bản ghi chưa tạo lợi thế rõ rệt cho mô hình phức tạp; ưu tiên mô hình đơn giúp giảm chi phí vận hành và kiểm thử.",
        "Phù hợp mục tiêu nguyên mẫu: đề tài cần chứng minh luồng tích hợp và tính minh bạch, không chỉ tối ưu một chỉ số dự báo.",
    ]:
        add_bullet(doc, item)
    add_body(doc, "Lựa chọn này không có nghĩa Decision Tree là mô hình tối ưu tuyệt đối. Nếu mục tiêu triển khai thay đổi sang ưu tiên sai số, Random Forest hoặc Gradient Boosting là ứng viên cần đánh giá tiếp bằng k-fold cross-validation, kiểm soát siêu tham số và dữ liệu nội bộ lớn hơn. Trong môi trường giáo dục, khả năng hiểu và kiểm tra dự báo là yêu cầu quan trọng đối với niềm tin và trách nhiệm giải trình [5], [6].")

    add_heading(doc, "CHƯƠNG 3. KIẾN TRÚC PHẦN MỀM VÀ KIẾN TRÚC AI", 1, page_break=True)
    add_heading(doc, "Kiến trúc phần mềm", 2)
    add_body(doc, "ScoreSense sử dụng kiến trúc client-service-model theo hướng tách giao diện khỏi xử lý AI. Flutter Web quản lý tương tác, thu thập dữ liệu và hiển thị kết quả; Flask cung cấp hợp đồng HTTP/JSON; pandas chuẩn hóa dữ liệu; artifact mô hình thực hiện suy luận. Kiến trúc này cho phép thay đổi giao diện hoặc mô hình độc lập trong giới hạn hợp đồng API.")
    add_picture(doc, software_diagram, 16.0)
    add_caption(doc, "Hình 1: Kiến trúc phần mềm ScoreSense xác minh từ mã nguồn và cấu hình triển khai")
    add_heading(doc, "Lớp giao diện Flutter Web", 3)
    add_body(doc, "Ứng dụng Flutter cung cấp hai luồng đầu vào. Luồng cá nhân sử dụng biểu mẫu tám bước và ánh xạ câu trả lời về đúng tên cột của mô hình. Luồng hàng loạt đọc tệp CSV, kiểm tra sự hiện diện của 32 cột bắt buộc rồi gửi toàn bộ ma trận gồm hàng tiêu đề và các hàng dữ liệu. Kết quả được hiển thị theo từng sinh viên và tổng hợp tỷ lệ Pass/Fail.")
    add_heading(doc, "Lớp dịch vụ Flask", 3)
    add_body(doc, "API được cài đặt bằng Flask 3.1.1, không phải FastAPI. Endpoint /predict kiểm tra cấu trúc JSON cơ bản, chuyển ma trận đầu vào thành DataFrame và gọi PredictScore. Flask-CORS chỉ cho phép các origin được cấu hình; MAX_CONTENT_LENGTH giới hạn request ở 10 MB. Khi triển khai, Gunicorn chạy một worker với timeout 120 giây trên Render.")
    add_heading(doc, "Kiến trúc triển khai", 3)
    add_body(doc, "Frontend được build với API_BASE_URL và triển khai trên Firebase Hosting; backend được triển khai độc lập trên Render theo render.yaml. Tách hai dịch vụ giúp frontend được phân phối như nội dung tĩnh, còn Python và artifact mô hình chỉ tồn tại ở backend. Hệ thống hiện không sử dụng cơ sở dữ liệu, không lưu lịch sử dự đoán và không có cơ chế đăng nhập/phân quyền.")
    add_heading(doc, "Kiến trúc AI", 2)
    add_body(doc, "Kiến trúc AI gồm hai nhánh. Nhánh huấn luyện hoạt động ngoại tuyến qua decisiontree.py, tạo tệp .pkl và danh sách cột .json. Nhánh suy luận hoạt động trực tuyến trong Flask; mô hình và schema được nạp theo version và cache tối đa ba phiên bản. Theo DEPLOYMENT.md, huấn luyện trực tuyến được chủ động vô hiệu hóa; vì vậy báo cáo không xem cập nhật mô hình là một API đang vận hành.")
    add_picture(doc, ai_diagram, 16.0)
    add_caption(doc, "Hình 2: Kiến trúc AI gồm huấn luyện ngoại tuyến và suy luận trực tuyến")
    add_heading(doc, "Luồng tiền xử lý và suy luận", 3)
    add_body(doc, "Dữ liệu mới được chuyển thành DataFrame; cột name được tách; biến phân loại được one-hot encoding với drop_first=True; DataFrame được reindex theo 41 cột huấn luyện và điền 0 cho cột thiếu. Mô hình được nạp theo phiên bản, trả điểm dự đoán, sau đó backend ghép name và originalIndex vào JSON. Giao diện dùng điều kiện prediction > 10 để gán Pass, ngược lại là Fail.")
    add_heading(doc, "Hợp đồng API", 2)
    add_caption(doc, "Bảng 4: Các endpoint của backend hiện tại")
    add_table(doc,
        ["Phương thức", "Endpoint", "Đầu vào", "Đầu ra / vai trò"],
        [
            ["GET", "/health", "Không", "{status: ok}; health check cho Render"],
            ["GET", "/totalmodel", "Không", "numModel; frontend dùng numModel - 1 để tạo danh sách version"],
            ["POST", "/predict", "version, inputDataImport", "Danh sách {name, prediction, originalIndex}"],
        ], [2.8, 3.4, 4.8, 5.0], font_size=9.8)
    add_heading(doc, "Yêu cầu phi chức năng và giới hạn kiến trúc", 2)
    add_body(doc, "Cấu hình CORS, health check, giới hạn kích thước request và timeout là các biện pháp vận hành đã có. Tuy nhiên, hệ thống chưa đáp ứng các yêu cầu triển khai với dữ liệu học vụ thật: xác thực, phân quyền, mã hóa dữ liệu nhạy cảm, nhật ký truy cập, lưu vết phiên bản mô hình, giám sát drift và cơ chế thu hồi kết quả. Những nội dung này phải được xem là công việc tương lai, không phải chức năng đã hoàn thành.")

    add_heading(doc, "CHƯƠNG 4. KẾT QUẢ VÀ KIỂM THỬ PHẦN MỀM", 1, page_break=True)
    add_heading(doc, "Kết quả tích hợp giao diện", 2)
    add_body(doc, "Sản phẩm đã hình thành các màn hình chính phục vụ lựa chọn phương thức dự đoán, nhập dữ liệu cá nhân, xem kết quả cá nhân và xem bảng kết quả hàng loạt. Bốn hình dưới đây là ảnh giao diện có trong tài liệu gốc và phù hợp với các route/module hiện có trong dự án.")
    for index, (image_path, caption) in enumerate(zip(screenshots, [
        "Giao diện lựa chọn phương thức dự đoán",
        "Giao diện biểu mẫu nhập dữ liệu cho một người học",
        "Giao diện kết quả dự đoán cá nhân",
        "Giao diện tổng hợp kết quả dự đoán hàng loạt",
    ]), start=3):
        add_picture(doc, image_path, 15.3)
        add_caption(doc, f"Hình {index}: {caption}")
    add_heading(doc, "Kết quả mô hình hiện tại", 2)
    add_body(doc, "Khi đánh giá lại artifact hiện tại trên tập kiểm thử 79 bản ghi được tạo bằng random_state=42, mô hình đạt MSE 4,2025; RMSE 2,0500; MAE 1,1392 và R² 0,7950. Các giá trị này thay thế nhóm số liệu MSE 6,33 và R² 0,69 trong bản thảo trước vì không tái lập được từ mô hình và dữ liệu đang lưu trong dự án.")
    add_caption(doc, "Bảng 5: Chỉ số đánh giá Decision Tree Regressor đang triển khai")
    add_table(doc,
        ["Chỉ số", "Giá trị", "Ý nghĩa"],
        [
            ["MSE", "4,2025", "Sai số bình phương trung bình"],
            ["RMSE", "2,0500", "Sai số quy về cùng đơn vị điểm G3"],
            ["MAE", "1,1392", "Độ lệch tuyệt đối trung bình"],
            ["R²", "0,7950", "Tỷ lệ biến thiên G3 được mô hình giải thích trên tập kiểm thử"],
        ], [3.3, 3.0, 9.7], font_size=10.5)
    add_heading(doc, "Phân tích vai trò đặc trưng", 2)
    add_body(doc, "Độ quan trọng theo mức giảm impurity cho thấy G2 chi phối 73,83%, tiếp theo là absences 15,46%. Điều này phù hợp với mô tả của bộ dữ liệu UCI rằng dự báo G3 trở nên dễ hơn đáng kể khi G1/G2 đã có [1]. Kết quả cũng cho thấy ScoreSense hiện gần với dự báo cuối kỳ hơn là cảnh báo thật sớm. Độ quan trọng không phải bằng chứng nhân quả và có thể thay đổi khi dữ liệu thay đổi.")
    add_caption(doc, "Bảng 6: Năm đặc trưng quan trọng nhất của mô hình hiện tại")
    add_table(doc,
        ["Đặc trưng", "Mức quan trọng", "Diễn giải thận trọng"],
        [
            ["G2", "73,83%", "Tín hiệu điểm gần G3 nhất"],
            ["absences", "15,46%", "Số buổi vắng học"],
            ["reason_home", "3,47%", "Một mức của lý do chọn trường sau one-hot"],
            ["age", "1,85%", "Tuổi trong phạm vi dữ liệu gốc"],
            ["famrel", "1,10%", "Mức chất lượng quan hệ gia đình tự khai báo"],
        ], [4.0, 3.5, 8.5], font_size=10.2)
    add_heading(doc, "Mục tiêu và phương pháp kiểm thử", 2)
    add_body(doc, "Kiểm thử được tổ chức theo ba lớp: kiểm thử hợp đồng backend bằng Flask test client; chạy bộ widget test có sẵn bằng flutter test; và phân tích tĩnh bằng flutter analyze. Mục tiêu là xác minh đường đi chính, phản hồi lỗi, tính nhất quán dữ liệu và chất lượng mã giao diện. Kết quả dưới đây phản ánh mã nguồn tại thời điểm rà soát, không phải số liệu giả định.")
    add_heading(doc, "Kiểm thử backend/API", 3)
    add_caption(doc, "Bảng 7: Ma trận kiểm thử backend đã thực thi")
    add_table(doc,
        ["Mã", "Tình huống", "Kỳ vọng", "Kết quả quan sát", "Đánh giá"],
        [
            ["BE‑01", "GET /health", "200, status=ok", "200, status=ok", "Đạt"],
            ["BE‑02", "GET /totalmodel", "Danh sách version dùng được", "200, numModel=4; UI sinh 3 version", "Đạt theo hợp đồng UI"],
            ["BE‑03", "Body không phải JSON object", "400", "400, thông báo rõ", "Đạt"],
            ["BE‑04", "inputDataImport thiếu header/hàng", "400", "400, thông báo rõ", "Đạt"],
            ["BE‑05", "version không phải số nguyên", "400", "400, thông báo rõ", "Đạt"],
            ["BE‑06", "version không tồn tại", "404", "404, thông báo rõ", "Đạt"],
            ["BE‑07", "Một bản ghi đủ 32 cột", "200, đúng schema", "200; name, prediction, originalIndex", "Đạt"],
            ["BE‑08", "Hai bản ghi có name", "Giữ tên và thứ tự", "Tên A/B và index 0/1 được giữ", "Đạt"],
            ["BE‑09", "Chỉ gửi G1 và G2", "Từ chối do thiếu cột", "200; backend điền 0 cho cột thiếu", "Chưa đạt"],
        ], [2.8, 3.3, 3.2, 4.2, 2.5], font_size=8.8)
    add_body(doc, "Kết quả cho thấy API xử lý tốt cấu trúc request cơ bản, phiên bản không hợp lệ và bảo toàn thứ tự khi dự đoán hàng loạt. Điểm chưa đạt là backend không kiểm tra đủ 32 trường bắt buộc; cơ chế reindex khiến request thiếu gần như toàn bộ đặc trưng vẫn nhận mã 200 và có thể tạo dự đoán không có ý nghĩa. Frontend có kiểm tra cột khi tải CSV, nhưng API vẫn cần tự bảo vệ vì có thể được gọi trực tiếp.")
    add_heading(doc, "Kiểm thử giao diện và phân tích tĩnh", 3)
    add_caption(doc, "Bảng 8: Kết quả kiểm thử Flutter hiện tại")
    add_table(doc,
        ["Hạng mục", "Kết quả", "Bằng chứng chính", "Kết luận"],
        [
            ["flutter test", "0 đạt, 2 thất bại", "RenderFlex tràn 340 px; test không tìm thấy Radio và nút Next", "Chưa đạt"],
            ["flutter analyze", "98 vấn đề", "Unused import/field, API deprecated, quy tắc đặt tên và cảnh báo khác", "Chưa đạt tiêu chí sạch cảnh báo"],
            ["Kiểm tra tệp CSV", "Có kiểm tra 32 cột", "requiredColumns.every(...) trước khi chuyển trang", "Đạt ở lớp UI"],
            ["Cấu hình gọi API", "Có", "API_BASE_URL qua --dart-define; timeout 120 giây", "Đạt"],
        ], [3.2, 3.1, 6.3, 3.4], font_size=9.2)
    add_body(doc, "Hai widget test hiện có chưa thể dùng làm bằng chứng giao diện ổn định vì đều thất bại trên kích thước mặc định của Flutter test. Lỗi tràn ngang cho thấy biểu mẫu chưa thích ứng đầy đủ với viewport hẹp; đồng thời test đang kỳ vọng widget Radio trong khi triển khai hiện tại đã thay đổi. Cần sửa cả bố cục và kịch bản kiểm thử trước khi công bố trạng thái đạt.")
    add_heading(doc, "Các rủi ro kỹ thuật được phát hiện", 2)
    for item in [
        "Quy tắc Pass dùng prediction > 10; điểm đúng bằng 10 bị xếp Fail. Ngưỡng cần được xác nhận và kiểm thử biên.",
        "Backend trả prediction dạng số thực, trong khi lớp Predictions của Flutter khai báo int; hợp đồng kiểu dữ liệu nên được chuẩn hóa để tránh lỗi giải mã.",
        "API chấp nhận dữ liệu thiếu cột và điền 0; cần validation theo schema và phạm vi giá trị.",
        "Cây sâu 11 với 127 lá có nguy cơ quá khớp; cần cắt tỉa và đánh giá chéo.",
        "Ba phiên bản mô hình chưa có metadata về dữ liệu, thời điểm huấn luyện, chỉ số và checksum; chưa đủ khả năng kiểm toán vòng đời mô hình.",
        "Chưa có xác thực, phân quyền, cơ sở dữ liệu, lưu lịch sử và nhật ký truy cập; không phù hợp để xử lý dữ liệu học vụ thật ở trạng thái hiện tại.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Đánh giá mức độ hoàn thành", 2)
    add_body(doc, "Đề tài đã hoàn thành nguyên mẫu đầu-cuối và chứng minh được khả năng tích hợp mô hình vào web. Phần backend cốt lõi hoạt động ở các tình huống hợp lệ và lỗi phổ biến. Tuy vậy, kết quả kiểm thử cho thấy sản phẩm đang ở mức prototype nghiên cứu: cần hoàn thiện validation, responsive UI, bộ test, quản trị mô hình và bảo mật trước khi thử nghiệm với người dùng thật.")

    add_heading(doc, "CHƯƠNG 5. KẾT LUẬN VÀ KIẾN NGHỊ", 1, page_break=True)
    add_heading(doc, "Kết luận", 2)
    add_body(doc, "ScoreSense đã tích hợp thành công Flutter Web, Flask API, quy trình pandas và Decision Tree Regressor thành một nguyên mẫu dự đoán điểm G3. Kiến trúc tách lớp giúp giao diện, dịch vụ và artifact mô hình có thể phát triển tương đối độc lập. Hai hình thức đầu vào - biểu mẫu cá nhân và CSV hàng loạt - đáp ứng đúng mục tiêu trình diễn của dự án.")
    add_body(doc, "So sánh thực nghiệm cho thấy Decision Tree là lựa chọn có cơ sở trong phạm vi hiện tại: MAE thấp nhất trong năm mô hình đối chứng, R² chỉ thấp hơn Random Forest khoảng 0,0198, artifact nhỏ và có khả năng truy vết điều kiện. Tuy nhiên, cây chưa được cắt tỉa và giao diện chưa cung cấp đường giải thích; vì vậy không nên mô tả hệ thống như một giải pháp XAI hoàn chỉnh.")
    add_body(doc, "Kết quả kiểm thử cũng làm rõ giới hạn sản phẩm. Backend đáp ứng phần lớn hợp đồng chính nhưng thiếu validation đặc trưng; hai widget test thất bại; phân tích tĩnh còn 98 vấn đề. Quan trọng hơn, mô hình được huấn luyện trên dữ liệu học sinh trung học Bồ Đào Nha, chưa phải dữ liệu sinh viên Văn Lang. Do đó, đóng góp hợp lý của đề tài là nguyên mẫu kỹ thuật và quy trình tích hợp, không phải mô hình đã sẵn sàng triển khai trong quản lý đào tạo.")
    add_heading(doc, "Kiến nghị theo mức ưu tiên", 2)
    priorities = [
        ("Ưu tiên 1 - Độ đúng và kiểm thử", "Bổ sung schema validation ở backend; thống nhất kiểu prediction; xác nhận ngưỡng Pass; sửa bố cục responsive và cập nhật widget test; thiết lập kiểm thử tự động cho API và Flutter."),
        ("Ưu tiên 2 - Chất lượng mô hình", "Dùng k-fold cross-validation; giới hạn max_depth/min_samples_leaf hoặc cắt tỉa cost-complexity; so sánh lại với Random Forest/Gradient Boosting; báo cáo khoảng tin cậy và sai số theo nhóm."),
        ("Ưu tiên 3 - Dữ liệu phù hợp bối cảnh", "Xây dựng bộ dữ liệu sinh viên Văn Lang có phê duyệt và đồng thuận; chuẩn hóa thang điểm, ngành học và học kỳ; đánh giá thiên lệch và khả năng tổng quát hóa trước triển khai."),
        ("Ưu tiên 4 - Kiến trúc vận hành", "Bổ sung xác thực, phân quyền, mã hóa, nhật ký truy cập, lưu lịch sử dự đoán và model registry có metadata/checksum; giám sát drift và quy trình quay lui phiên bản."),
        ("Ưu tiên 5 - Khả năng giải thích", "Hiển thị đường đi của mẫu qua cây, các điều kiện đóng góp và cảnh báo về giới hạn; chỉ chuyển đặc trưng có thể can thiệp thành gợi ý học tập, tránh diễn giải nhân quả từ feature importance."),
    ]
    for lead, content in priorities:
        add_body(doc, f"{lead}: {content}", first_line=False, bold_lead=f"{lead}:")
    add_heading(doc, "Nguyên tắc sử dụng có trách nhiệm", 2)
    add_body(doc, "Khi mở rộng sang dữ liệu thật, kết quả dự đoán chỉ nên là tín hiệu hỗ trợ, không phải căn cứ duy nhất cho quyết định ảnh hưởng đến quyền lợi người học. Sinh viên cần được thông báo mục đích sử dụng dữ liệu, quyền truy cập và cơ chế phản hồi. Việc áp dụng cần tuân thủ nguyên tắc lấy con người làm trung tâm, minh bạch, công bằng và có trách nhiệm giải trình [7], [8].")

    add_title(doc, "TÀI LIỆU THAM KHẢO")
    references = [
        "[1] Cortez, P. và Silva, A. M. G. (2008), 'Using Data Mining to Predict Secondary School Student Performance', Proceedings of the 5th Annual Future Business Technology Conference, tr. 5-12.",
        "[2] Cortez, P. (2008), Student Performance [Dataset], UCI Machine Learning Repository, DOI: 10.24432/C5TG7T, https://archive.ics.uci.edu/dataset/320/student+performance.",
        "[3] Breiman, L., Friedman, J. H., Olshen, R. A. và Stone, C. J. (1984), Classification and Regression Trees, Wadsworth.",
        "[4] Breiman, L. (2001), 'Random Forests', Machine Learning, 45, tr. 5-32, DOI: 10.1023/A:1010933404324.",
        "[5] Alyahyan, E. và Düştegör, D. (2020), 'Predicting academic success in higher education: literature review and best practices', International Journal of Educational Technology in Higher Education, 17, bài 3, DOI: 10.1186/s41239-020-0177-7.",
        "[6] Lünich, M. (2024), 'Explainable Artificial Intelligence for Academic Performance Prediction', Proceedings of the 2024 ACM Conference on Fairness, Accountability, and Transparency, DOI: 10.1145/3630106.3658953.",
        "[7] Thủ tướng Chính phủ (2022), Quyết định số 131/QĐ-TTg về tăng cường ứng dụng công nghệ thông tin và chuyển đổi số trong giáo dục và đào tạo giai đoạn 2022-2025, định hướng đến năm 2030.",
        "[8] UNESCO (2021), Recommendation on the Ethics of Artificial Intelligence, Paris: UNESCO.",
    ]
    for ref in references:
        p = doc.add_paragraph(style="List Paragraph")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(ref)
        format_run(r, size=11.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Ensure fields (TOC and table/figure lists) refresh when opened.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.core_properties.title = "Báo cáo tổng kết đề tài ScoreSense - bản bổ sung kiến trúc và kiểm thử"
    doc.core_properties.subject = "Kiến trúc phần mềm, kiến trúc AI, so sánh mô hình và kiểm thử"
    finalize_accessibility(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
