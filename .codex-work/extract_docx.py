from pathlib import Path
from docx import Document


source = Path(r"F:\Downloads\BM-29-NCKHSV-bao-cao-tong-ket-de-tai.docx")
output = Path(r"F:\Github\Do_An_Ung_Dung_2\.codex-work\report_extract.txt")
doc = Document(source)

lines = [
    f"PARAGRAPHS {len(doc.paragraphs)} TABLES {len(doc.tables)} "
    f"SECTIONS {len(doc.sections)} INLINE_SHAPES {len(doc.inline_shapes)}"
]
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.replace("\n", " | ").strip()
    if text:
        lines.append(f"P{i:04d} [{paragraph.style.name}] {text}")

for table_index, table in enumerate(doc.tables):
    lines.append(f"---TABLE {table_index} {len(table.rows)}x{len(table.columns)}---")
    for row_index, row in enumerate(table.rows):
        cell_text = " || ".join(cell.text.replace("\n", " | ") for cell in row.cells)
        lines.append(f"R{row_index}: {cell_text}")

output.parent.mkdir(parents=True, exist_ok=True)
output.write_text("\n".join(lines), encoding="utf-8")
